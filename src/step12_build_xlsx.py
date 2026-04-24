"""
Step 12: Build a downloadable Excel workbook for the website.

Sheets:
  1. Habilidades  — one row per habilidade with top CC match
  2. Componentes  — one row per BNCC component (PT + EN + source tag + tier)
  3. Pré-requisitos — one row per prerequisite edge

Output: docs/data/bncc_learning_commons.xlsx
"""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT     = Path(__file__).resolve().parent.parent
PROC     = ROOT / "data" / "processed"
RAW      = ROOT / "data" / "raw"
OUT_XLSX = ROOT / "docs" / "data" / "bncc_learning_commons.xlsx"
LOGS_DIR = ROOT / "data" / "logs"
LOGS_DIR.mkdir(parents=True, exist_ok=True)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s: %(message)s",
    handlers=[
        logging.FileHandler(LOGS_DIR / "step12_build_xlsx.log", mode="w"),
        logging.StreamHandler(),
    ],
)
log = logging.getLogger(__name__)

# ── Styles ───────────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="16172A")
HEADER_FONT = Font(bold=True, color="D4D6F0", size=10)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="2A2B45")
BORDER = Border(bottom=THIN)

LC_FILL  = PatternFill("solid", fgColor="1D3461")
AI_FILL  = PatternFill("solid", fgColor="3A2800")
LC_FONT  = Font(color="60A5FA", bold=True, size=9)
AI_FONT  = Font(color="FBBF24", bold=True, size=9)


def style_header(ws, headers: list[str], col_widths: list[int]) -> None:
    ws.append(headers)
    for col_idx, (_, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill   = HEADER_FILL
        cell.font   = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def build_habilidades_sheet(ws, bncc, matches_s, cc_stds) -> None:
    headers = [
        "Código", "Etapa", "Série", "Disciplina", "Descrição (PT)",
        "Código CC (melhor match)", "Descrição CC", "Score CC",
    ]
    widths  = [12, 8, 10, 28, 70, 20, 60, 10]
    style_header(ws, headers, widths)

    top_match = (
        matches_s[matches_s["rank"] == 1]
        .set_index("habilidade_code")[["cc_standard_id", "score"]]
        .to_dict("index")
    )
    cc_code_map = dict(zip(cc_stds["identifier"], cc_stds["statement_code"]))
    cc_desc_map = dict(zip(cc_stds["identifier"], cc_stds["description"]))

    for _, row in bncc.iterrows():
        code = row["bncc_code"]
        tm   = top_match.get(code, {})
        sid  = tm.get("cc_standard_id", "")
        ws.append([
            code,
            row.get("stage", ""),
            row.get("grade", ""),
            row.get("subject", ""),
            row.get("description_pt", ""),
            cc_code_map.get(sid, ""),
            cc_desc_map.get(sid, ""),
            round(float(tm["score"]), 4) if tm.get("score") else "",
        ])

    log.info("Habilidades sheet: %d rows", ws.max_row - 1)


def build_componentes_sheet(ws, comps_pt, comps_en, matches_c) -> None:
    headers = [
        "Código Habilidade", "ID Componente", "Descrição (PT)", "Descrição (EN)",
        "Fonte", "Tier de correspondência", "Componente CC correspondente",
    ]
    widths  = [18, 38, 60, 60, 8, 20, 60]
    style_header(ws, headers, widths)

    match_map = {
        row["bncc_component_id"]: row["match_tier"]
        for _, row in matches_c.iterrows()
    }
    cc_comp_en = (
        pd.read_csv(ROOT / "data" / "processed" / "cc_components_full.csv", encoding="utf-8-sig")
        .set_index("component_id")["description"].to_dict()
    )

    en_map = (
        comps_en.drop_duplicates("component_id")
        .set_index("component_id")
    )
    pt_map = (
        comps_pt.drop_duplicates("component_id")
        .set_index("component_id")["description_pt"].to_dict()
    )

    row_num = 2
    for _, row in comps_en.iterrows():
        cid   = str(row["component_id"])
        tier  = match_map.get(cid, "none")
        source = "LC" if tier == "merge" else "AI"

        # For merge components, the component_id IS the CC component_id
        cc_desc = cc_comp_en.get(cid, "") if tier == "merge" else ""

        ws.append([
            row["habilidade_code"],
            cid,
            pt_map.get(cid, ""),
            row.get("description", ""),
            source,
            tier,
            cc_desc,
        ])

        # Colour the source cell
        source_cell = ws.cell(row=row_num, column=5)
        if source == "LC":
            source_cell.fill = LC_FILL
            source_cell.font = LC_FONT
        else:
            source_cell.fill = AI_FILL
            source_cell.font = AI_FONT
        source_cell.alignment = Alignment(horizontal="center")
        row_num += 1

    log.info("Componentes sheet: %d rows", ws.max_row - 1)


def build_prereqs_sheet(ws, prereqs, bncc) -> None:
    headers = [
        "Código Habilidade", "Descrição Habilidade (PT)",
        "Código Pré-requisito", "Descrição Pré-requisito (PT)",
        "Padrão CC (fonte)", "Score hab. CC", "Score pré. CC",
    ]
    widths = [18, 65, 18, 65, 20, 12, 12]
    style_header(ws, headers, widths)

    desc_map = dict(zip(bncc["bncc_code"], bncc["description_pt"]))

    for _, row in prereqs.iterrows():
        ws.append([
            row["habilidade_code"],
            desc_map.get(row["habilidade_code"], ""),
            row["prerequisite_code"],
            desc_map.get(row["prerequisite_code"], ""),
            row.get("source_cc_prereq_code", ""),
            round(float(row["best_hab_cc_score"]), 4),
            round(float(row["best_pre_cc_score"]), 4),
        ])

    log.info("Pré-requisitos sheet: %d rows", ws.max_row - 1)


def main() -> None:
    log.info("=== Step 12: Build XLSX workbook ===")

    bncc      = pd.read_csv(PROC / "bncc_translated.csv",          encoding="utf-8-sig")
    comps_en  = pd.read_csv(PROC / "bncc_components.csv",          encoding="utf-8-sig")
    comps_pt  = pd.read_csv(PROC / "bncc_components_pt.csv",       encoding="utf-8-sig")
    matches_c = pd.read_csv(PROC / "bncc_component_matches.csv",   encoding="utf-8-sig")
    matches_s = pd.read_csv(PROC / "bncc_cc_standard_matches.csv", encoding="utf-8-sig")
    prereqs   = pd.read_csv(PROC / "bncc_prerequisites.csv",       encoding="utf-8-sig")
    cc_stds   = pd.read_csv(RAW  / "lc_cc_standards.csv",          encoding="utf-8-sig")

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Habilidades"
    build_habilidades_sheet(ws1, bncc, matches_s, cc_stds)

    ws2 = wb.create_sheet("Componentes")
    build_componentes_sheet(ws2, comps_pt, comps_en, matches_c)

    ws3 = wb.create_sheet("Pré-requisitos")
    build_prereqs_sheet(ws3, prereqs, bncc)

    wb.save(OUT_XLSX)
    size_kb = OUT_XLSX.stat().st_size // 1024
    log.info("Saved %s (%d KB)", OUT_XLSX.name, size_kb)


if __name__ == "__main__":
    main()
